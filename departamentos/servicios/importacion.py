"""Importación de departamentos y lecturas desde Excel."""
import unicodedata

from django.db import transaction

from ..models import Departamento
from .agua import registrar_lectura
from .comun import a_decimal

# Encabezados aceptados (normalizados) -> campo interno.
_ENCABEZADOS = {
    'DEPARTAMENTO': 'numero', 'DPTO': 'numero', 'DEPTO': 'numero',
    'NUMERO': 'numero', 'N_DEPARTAMENTO': 'numero', 'N_DPTO': 'numero',
    'LECTURA_ANTERIOR': 'anterior', 'ANTERIOR': 'anterior',
    'LECTURA_ACTUAL': 'actual', 'ACTUAL': 'actual',
    'COCHERA': 'cochera', 'MONTO_COCHERA': 'cochera', 'N_COCHERA': 'cochera',
}


def _norm(valor):
    """Normaliza un encabezado: mayúsculas, sin acentos, con guion bajo."""
    texto = str(valor or '').strip().upper()
    texto = ''.join(c for c in unicodedata.normalize('NFD', texto)
                    if unicodedata.category(c) != 'Mn')
    return texto.replace(' ', '_')


def _buscar_encabezados(ws):
    """(fila, {campo: columna}) de la fila de encabezados (busca en las primeras filas)."""
    for r in range(1, min(ws.max_row, 15) + 1):
        fila = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            campo = _ENCABEZADOS.get(_norm(ws.cell(row=r, column=c).value))
            if campo:
                fila[campo] = c
        if 'numero' in fila:
            return r, fila
    raise ValueError('No se encontró la columna "DEPARTAMENTO" en la primera hoja.')


@transaction.atomic
def importar_lecturas_excel(fileobj, periodo):
    """Lee un Excel con el formato del sistema y carga departamentos + lecturas.

    Formato esperado (primera hoja): una fila de encabezados con al menos la
    columna DEPARTAMENTO, y opcionalmente LECTURA_ANTERIOR, LECTURA_ACTUAL,
    COCHERA. Una fila por departamento.
    """
    import openpyxl

    ws = openpyxl.load_workbook(fileobj, data_only=True).active
    header_row, colmap = _buscar_encabezados(ws)

    def celda(r, campo):
        return a_decimal(ws.cell(row=r, column=colmap[campo]).value) if campo in colmap else None

    resumen = {'departamentos': 0, 'lecturas': 0, 'errores': []}
    for r in range(header_row + 1, ws.max_row + 1):
        crudo = ws.cell(row=r, column=colmap['numero']).value
        if crudo is None or str(crudo).strip() == '':
            continue
        # Normaliza el número (101.0 -> "101"; "T6-301" -> "T6-301").
        num_dec = a_decimal(crudo)
        numero = str(int(num_dec)) if (num_dec is not None and num_dec == num_dec.to_integral_value()) \
            else str(crudo).strip()

        defaults = {'activo': True}
        cochera = celda(r, 'cochera')
        if cochera is not None:
            defaults['cochera_monto'] = cochera
        dep, _ = Departamento.objects.update_or_create(numero_domicilio=numero, defaults=defaults)
        resumen['departamentos'] += 1

        actual = celda(r, 'actual')
        if actual is not None:
            try:
                registrar_lectura(dep, periodo, lectura_actual=actual, lectura_anterior=celda(r, 'anterior'))
                resumen['lecturas'] += 1
            except Exception as e:  # pragma: no cover - fila con datos inconsistentes
                resumen['errores'].append(f'Depto {numero}: {e}')

    return resumen
