"""Carga inicial de deudas (migración desde el Excel del condominio).

Dos partes:
- `leer_excel_deudas`: parser PURO del libro (hojas DEUDAS ACTUALES y DEUDAS
  ANTERIORES). Devuelve filas con la misma forma que la hoja, el total de control,
  errores y advertencias. No toca la base.
- `deuda_migrada` / `guardar_deuda_migrada`: llevan esas filas a cargos reales
  (`Cargo.origen = MIGRACION`) y los vuelven a leer en la misma forma, para
  editarlos en la grilla hasta que cuadre.

Dónde queda cada monto (con el mes de corte, p. ej. agosto 2026):
- Meses de mantenimiento (JULIO, AGOSTO, y el detalle de D. ANTERIORES) →
  Mantenimiento en su mes real.
- Montos sin mes (D. ANTERIORES sin detalle, filas sin año) → "Deuda anterior"
  en el mes de corte.
- Puertas, Cámaras, Ascensores, Cuota, Multa, Mora → su concepto en el mes de corte.
"""
import unicodedata
from collections import defaultdict
from datetime import date
from decimal import Decimal, InvalidOperation

from django.db import transaction

from ..models import Cargo, Concepto, Imputacion, Pago
from .comun import CERO, fecha_vencimiento, redondear

DEUDA_ANT = 'DEUDA_ANT'
# Columnas de "servicios" y "otros" de DEUDAS ACTUALES → código de concepto.
SERVICIOS = [('puertas', 'PUERTAS', 'Puertas'), ('camara', 'CAMARA', 'Cámaras'),
             ('ascensores', 'ASCENSORES', 'Ascensores')]
OTROS = [('cuota', 'CUOTA', 'Cuota'), ('multa', 'MULTA', 'Multa'), ('mora', 'MORA', 'Mora')]
CAMPOS_CONCEPTO = SERVICIOS + OTROS

_MESES = {'ENE': 1, 'FEB': 2, 'MAR': 3, 'ABR': 4, 'MAY': 5, 'JUN': 6, 'JUL': 7,
          'AGO': 8, 'SET': 9, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DIC': 12}
_ENCABEZADOS = {'PUERTAS': 'puertas', 'PUERTA': 'puertas', 'CAMARA': 'camara', 'CAMARAS': 'camara',
                'ASCENSORES': 'ascensores', 'ASCENSOR': 'ascensores',
                'D ANTERIORES': 'd_ant', 'D ANTERIOR': 'd_ant', 'DEUDAS ANTERIORES': 'd_ant',
                'CUOTA': 'cuota', 'CUOTAS': 'cuota', 'MULTA': 'multa', 'MULTAS': 'multa',
                'MORA': 'mora', 'MORAS': 'mora', 'TOTAL': 'total'}


def _norm(v):
    """'D. ANTERIORES ' → 'D ANTERIORES'; 'Cámara' → 'CAMARA'."""
    t = ''.join(c for c in unicodedata.normalize('NFD', str(v or '').upper())
                if unicodedata.category(c) != 'Mn')
    return ' '.join(t.replace('.', ' ').replace('#', ' ').split())


def mes_de(encabezado):
    """Número de mes si el encabezado es un mes ('JULIO', 'AGOS.', 'SET.'), si no None."""
    t = _norm(encabezado)
    return _MESES.get(t[:3]) if len(t) >= 3 and t.isalpha() else None


def _monto(v):
    """(Decimal o None, es_texto). Vacío → None; número → 2 decimales (quita ruido de float)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, False
    if isinstance(v, bool):
        return None, True
    if isinstance(v, (int, float)):
        return redondear(Decimal(str(v))), False
    try:
        return redondear(Decimal(str(v).strip().replace(',', ''))), False
    except InvalidOperation:
        return None, True


def _dpto(v):
    if v is None or str(v).strip() == '':
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _hoja(wb, nombre):
    """Hoja por nombre sin distinguir mayúsculas ni espacios de más."""
    buscado = _norm(nombre)
    return next((ws for ws in wb.worksheets if _norm(ws.title) == buscado), None)


def _buscar_encabezado(filas, clave, max_filas=15):
    """(índice de fila, fila) de la primera fila que tiene una celda == `clave` (normalizada)."""
    for i, fila in enumerate(filas[:max_filas]):
        if any(_norm(c) == clave for c in fila):
            return i, fila
    return None, None


def _leer_actuales(ws, anio_meses):
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    i, cab = _buscar_encabezado(filas, 'DPTO')
    if cab is None:
        return None, ['DEUDAS ACTUALES: no se encontró la fila con "# DPTO".']
    col_dpto = next(k for k, c in enumerate(cab) if _norm(c) == 'DPTO')
    columnas = {}   # índice → campo ('puertas', 'mes:7', 'total', ...)
    for k, c in enumerate(cab):
        if k <= col_dpto or c is None:
            continue
        m = mes_de(c)
        campo = f'mes:{m}' if m else _ENCABEZADOS.get(_norm(c))
        if campo:
            columnas[k] = campo
    meses = sorted(int(c[4:]) for c in columnas.values() if c.startswith('mes:'))

    resultado, total_control, errores = [], None, []
    for n, fila in enumerate(filas[i + 1:], start=i + 2):
        dpto = _dpto(fila[col_dpto] if col_dpto < len(fila) else None)
        valores, textos = {}, []
        for k, campo in columnas.items():
            v, es_texto = _monto(fila[k] if k < len(fila) else None)
            if es_texto:
                textos.append(campo)
            valores[campo] = v
        if not dpto:
            if valores.get('total') is not None and total_control is None:
                total_control = valores['total']
            continue
        conceptos = {campo: valores.get(campo) or CERO for campo, _, _ in CAMPOS_CONCEPTO}
        r = {
            'dpto': dpto, 'fila_excel': n,
            **conceptos,
            'd_ant': valores.get('d_ant') or CERO,
            'meses': {f'{anio_meses}-{m:02d}': valores.get(f'mes:{m}') or CERO for m in meses},
            'total_excel': valores.get('total'),
            'errores': [], 'advertencias': [f'Celda con texto ignorada ({c}).' for c in textos],
        }
        suma = sum(conceptos.values(), CERO) + r['d_ant'] + sum(r['meses'].values(), CERO)
        if r['total_excel'] is not None and r['total_excel'] != suma:
            r['advertencias'].append(f'TOTAL del Excel {r["total_excel"]:.2f} ≠ suma de columnas {suma:.2f}.')
        resultado.append(r)
    return {'filas': resultado, 'meses': [f'{anio_meses}-{m:02d}' for m in meses],
            'total_control': total_control}, errores


def _leer_anteriores(ws):
    """{dpto: {'meses': {'2022-01': monto}, 'sin_periodo': monto, 'puertas': monto, 'advertencias': []}}."""
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    i, cab = _buscar_encabezado(filas, 'DEPART')
    if cab is None:
        return {}, ['DEUDAS ANTERIORES: no se encontró la fila con "DEPART.".']
    col = {'dpto': next(k for k, c in enumerate(cab) if _norm(c) == 'DEPART')}
    meses, fin = {}, len(cab)
    for k, c in enumerate(cab):
        t = _norm(c)
        if t in ('ANO', 'AÑO'):
            col['anio'] = k
        elif t == 'PUERTAS':
            col['puertas'] = k
        elif t == 'TOTAL' and 'total' not in col:
            col['total'] = fin = k           # lo de la derecha de TOTAL es otra tabla: se ignora
        elif mes_de(c) and 'total' not in col and k > col['dpto']:
            meses[k] = mes_de(c)

    detalle, actual = {}, None
    for n, fila in enumerate(filas[i + 1:], start=i + 2):
        fila = fila[:fin + 1]
        etiqueta = ' '.join(_norm(c) for c in fila if isinstance(c, str))
        if 'TOTAL GENERAL' in etiqueta:
            break
        dpto = _dpto(fila[col['dpto']] if col['dpto'] < len(fila) else None)
        if dpto:
            actual = detalle.setdefault(dpto, {'meses': {}, 'sin_periodo': CERO, 'puertas': CERO,
                                               'advertencias': []})
        if actual is None:
            continue
        anio = fila[col['anio']] if 'anio' in col and col['anio'] < len(fila) else None
        try:
            anio = int(anio) if anio not in (None, '') else None
        except (TypeError, ValueError):
            anio = None
        hay_monto = False
        for k, m in meses.items():
            v, es_texto = _monto(fila[k] if k < len(fila) else None)
            if es_texto:
                actual['advertencias'].append(
                    f'Fila {n}: texto en {m:02d}/{anio or "?"} tratado como vacío.')
            if v is None:
                continue
            hay_monto = True
            if anio:
                clave = f'{anio}-{m:02d}'
                actual['meses'][clave] = actual['meses'].get(clave, CERO) + v
            else:
                actual['sin_periodo'] += v
        if 'puertas' in col:
            v, _ = _monto(fila[col['puertas']] if col['puertas'] < len(fila) else None)
            if v:
                actual['puertas'] += v
                hay_monto = True
        if not hay_monto:
            continue   # subtotal (solo TOTAL) o fila vacía
    return detalle, []


def _leer_cobradas(ws, anio):
    """Tabla de la derecha de DEUDAS ANTERIORES ("# DEPT", MAYO, JUNIO, TOTAL): deudas
    de esos meses que ya se cobraron. Devuelve ({dpto: {'meses', 'total_excel',
    'advertencias'}}, total de control)."""
    filas = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, cab in enumerate(filas[:15]):
        col_dpto = next((k for k, c in enumerate(cab) if _norm(c) == 'DEPT'), None)
        if col_dpto is not None:
            break
    else:
        return {}, None
    meses, col_total = {}, None
    for k in range(col_dpto + 1, len(cab)):
        if mes_de(cab[k]):
            meses[k] = mes_de(cab[k])
        elif _norm(cab[k]) == 'TOTAL':
            col_total = k
            break
    cobradas, control = {}, None
    celda = lambda fila, k: fila[k] if k is not None and k < len(fila) else None
    for n, fila in enumerate(filas[i + 1:], start=i + 2):
        dpto = _dpto(celda(fila, col_dpto))
        total, _ = _monto(celda(fila, col_total))
        if not dpto:
            if total is not None and control is None and cobradas:
                control = total
            continue
        r = cobradas.setdefault(dpto, {'meses': {}, 'total_excel': total, 'advertencias': []})
        for k, m in meses.items():
            v, es_texto = _monto(celda(fila, k))
            if es_texto:
                r['advertencias'].append(f'Deudas cobradas, fila {n}: texto en {m:02d}/{anio} tratado como vacío.')
            if v:
                clave = f'{anio}-{m:02d}'
                r['meses'][clave] = r['meses'].get(clave, CERO) + v
        suma = sum(r['meses'].values(), CERO)
        if total is not None and total != suma:
            r['advertencias'].append(f'Deudas cobradas: TOTAL {total:.2f} ≠ suma de meses {suma:.2f}.')
    return cobradas, control


def leer_excel_deudas(fileobj, anio_meses):
    """Parser puro del Excel de deudas. `anio_meses` = año de las columnas de mes
    de DEUDAS ACTUALES (JULIO, AGOSTO...).

    Devuelve {'hojas': [(nombre, reconocida)], 'meses': ['2026-07', ...],
    'filas': [...], 'total_control', 'total_calculado', 'cobradas': {dpto: {...}},
    'cobradas_control', 'cobradas_total', 'errores', 'advertencias'}.
    `cobradas` = tabla de deudas de mayo/junio ya cobradas (van como ingreso).
    Cada fila: dpto, puertas, camara, ascensores, d_ant, meses{}, cuota, multa,
    mora, total, detalle{'meses', 'sin_periodo'}, errores[], advertencias[].
    """
    import openpyxl

    wb = openpyxl.load_workbook(fileobj, read_only=True, data_only=True)
    act_ws, ant_ws = _hoja(wb, 'DEUDAS ACTUALES'), _hoja(wb, 'DEUDAS ANTERIORES')
    out = {'hojas': [(ws.title, ws in (act_ws, ant_ws)) for ws in wb.worksheets],
           'meses': [], 'filas': [], 'total_control': None, 'total_calculado': CERO,
           'cobradas': {}, 'cobradas_control': None, 'cobradas_total': CERO,
           'errores': [], 'advertencias': []}
    if act_ws is None:
        out['errores'].append('No se encontró la hoja "DEUDAS ACTUALES".')
        return out
    actuales, errs = _leer_actuales(act_ws, anio_meses)
    out['errores'] += errs
    if actuales is None:
        return out
    detalle, errs = _leer_anteriores(ant_ws) if ant_ws is not None else ({}, [])
    out['errores'] += errs
    if ant_ws is None:
        out['advertencias'].append('No se encontró la hoja "DEUDAS ANTERIORES": '
                                   'D. ANTERIORES irá como un solo monto por depto.')

    vistos = set()
    for r in actuales['filas']:
        if r['dpto'] in vistos:
            r['errores'].append('Depto repetido en el Excel.')
        vistos.add(r['dpto'])
        det = detalle.get(r['dpto'])
        if det:
            r['advertencias'] += det['advertencias']
        r['detalle'] = conciliar(r['d_ant'], det)
        if r['detalle'] is None:
            suma = sum(det['meses'].values(), CERO) + det['sin_periodo']
            r['errores'].append(f'D. ANTERIORES {r["d_ant"]:.2f} ≠ detalle mes a mes {suma:.2f} '
                                '(sin contar puertas). Corrige uno de los dos.')
            r['detalle'] = {'meses': dict(det['meses']), 'sin_periodo': det['sin_periodo']}
        r['total'] = total_fila(r)
        out['total_calculado'] += r['total']
    out['filas'], out['meses'], out['total_control'] = actuales['filas'], actuales['meses'], actuales['total_control']
    if ant_ws is not None:
        out['cobradas'], out['cobradas_control'] = _leer_cobradas(ant_ws, anio_meses)
        out['cobradas_total'] = sum((sum(c['meses'].values(), CERO) for c in out['cobradas'].values()), CERO)
        if out['cobradas_control'] is not None and out['cobradas_control'] != out['cobradas_total']:
            out['advertencias'].append(f'Deudas cobradas: el TOTAL de la tabla ({out["cobradas_control"]:.2f}) '
                                       f'no coincide con la suma de las filas ({out["cobradas_total"]:.2f}).')
    if out['total_control'] is not None and out['total_control'] != out['total_calculado']:
        faltan = [r['dpto'] for r in out['filas'] if r['total'] == out['total_calculado'] - out['total_control']]
        out['advertencias'].append(
            f'El TOTAL de la hoja ({out["total_control"]:.2f}) no coincide con la suma de las filas '
            f'({out["total_calculado"]:.2f})' + (f'; la fórmula parece omitir al depto {", ".join(faltan)}.' if faltan else '.'))
    return out


def conciliar(d_ant, det):
    """Detalle de D. ANTERIORES: el mes a mes si cuadra; si no hay detalle, un solo
    monto sin periodo; None si el detalle no cuadra (error)."""
    if not det or (not det['meses'] and not det['sin_periodo']):
        return {'meses': {}, 'sin_periodo': d_ant}
    suma = sum(det['meses'].values(), CERO) + det['sin_periodo']
    if suma != d_ant:
        return None
    return {'meses': dict(sorted(det['meses'].items())), 'sin_periodo': det['sin_periodo']}


def total_fila(r):
    return (sum((r[c] for c, _, _ in CAMPOS_CONCEPTO), CERO) + r['d_ant']
            + sum(r['meses'].values(), CERO))


# ------------------------------------------------------------ Cargos migrados

def _periodo(clave):
    anio, mes = clave.split('-')
    return date(int(anio), int(mes), 1)


def meses_columna(corte, n=2):
    """Las `n` columnas de mes de la grilla: los últimos meses hasta el corte (jul, ago)."""
    out, y, m = [], corte.year, corte.month
    for _ in range(n):
        out.append(f'{y}-{m:02d}')
        y, m = (y - 1, 12) if m == 1 else (y, m - 1)
    return sorted(out)


def _concepto(codigo):
    concepto = Concepto.objects.filter(codigo=codigo).first()
    if concepto is None:   # base sin los conceptos estándar (ej. "Deuda anterior" nuevo)
        from .cuenta import sembrar_conceptos
        sembrar_conceptos()
        concepto = Concepto.objects.get(codigo=codigo)
    return concepto


def cobrado_migracion(cargo):
    """Parte del cargo saldada por pagos de la carga inicial (deuda ya cobrada)."""
    return sum((i.monto for i in cargo.imputaciones.all() if i.pago.origen == Cargo.Origen.MIGRACION), CERO)


def deuda_migrada(corte, meses):
    """Cargos de migración leídos en la forma de la grilla: {dpto: fila}.
    Mantenimiento de los meses de `meses` va a su columna; el anterior (y la
    "Deuda anterior" sin mes) forma D. ANTERIORES con su detalle. Lo saldado por
    pagos de la carga inicial va aparte, en `cobradas` (deuda ya cobrada)."""
    filas = {}
    cargos = (Cargo.objects.filter(origen=Cargo.Origen.MIGRACION)
              .select_related('departamento', 'concepto').prefetch_related('imputaciones__pago'))
    for c in cargos:
        dep = c.departamento.numero_domicilio
        r = filas.setdefault(dep, fila_vacia(dep, meses))
        cod, clave = c.concepto.codigo, f'{c.periodo:%Y-%m}'
        cobrado = cobrado_migracion(c) if cod == Concepto.MANTENIMIENTO else CERO
        pendiente = c.monto - cobrado
        if cobrado:
            r['cobradas'][clave] = r['cobradas'].get(clave, CERO) + cobrado
        campo = next((f for f, k, _ in CAMPOS_CONCEPTO if k == cod), None)
        if campo:
            r[campo] += pendiente
        elif cod == Concepto.MANTENIMIENTO and clave in r['meses']:
            r['meses'][clave] += pendiente
        elif cod == Concepto.MANTENIMIENTO and pendiente:
            r['detalle']['meses'][clave] = r['detalle']['meses'].get(clave, CERO) + pendiente
        elif cod != Concepto.MANTENIMIENTO:   # DEUDA_ANT u otro concepto: deuda anterior sin mes
            r['detalle']['sin_periodo'] += pendiente
        r['pagado'] += c.total_imputado - cobrado
    for r in filas.values():
        r['detalle']['meses'] = dict(sorted(r['detalle']['meses'].items()))
        r['cobradas'] = dict(sorted(r['cobradas'].items()))
        r['d_ant'] = sum(r['detalle']['meses'].values(), CERO) + r['detalle']['sin_periodo']
        r['total'] = total_fila(r)
    return filas


def fila_vacia(dpto, meses):
    return {'dpto': dpto, **{c: CERO for c, _, _ in CAMPOS_CONCEPTO}, 'd_ant': CERO,
            'meses': {m: CERO for m in meses}, 'detalle': {'meses': {}, 'sin_periodo': CERO},
            'cobradas': {}, 'total': CERO, 'pagado': CERO, 'errores': [], 'advertencias': []}


def cargos_deseados(r, corte):
    """{(codigo_concepto, periodo): monto} que debe tener el depto según la fila
    (lo pendiente + lo ya cobrado de ese mes, que se salda con su pago)."""
    out = defaultdict(lambda: CERO)
    for campo, cod, _ in CAMPOS_CONCEPTO:
        if r[campo]:
            out[(cod, corte)] += r[campo]
    meses = list(r['meses'].items()) + list(r['detalle']['meses'].items()) + list(r.get('cobradas', {}).items())
    for clave, v in meses:
        if v:
            out[(Concepto.MANTENIMIENTO, _periodo(clave))] += v
    if r['detalle']['sin_periodo']:
        out[(DEUDA_ANT, corte)] += r['detalle']['sin_periodo']
    return {k: redondear(v) for k, v in out.items() if v}


@transaction.atomic
def guardar_deuda_migrada(dep, r, corte, fecha_cobro=None):
    """Deja los cargos de migración del depto iguales a la fila `r` y registra lo ya
    cobrado (`r['cobradas']`) como pagos de la carga inicial con fecha `fecha_cobro`
    (cuentan como ingreso de ese mes en el balance). Devuelve (creados, actualizados,
    borrados) o lanza ValueError con el motivo (nada se guarda para ese depto). No
    toca cargos del sistema ni baja un cargo por debajo de lo ya pagado."""
    deseados = cargos_deseados(r, corte)
    for (cod, periodo), _ in deseados.items():
        if periodo > corte:
            raise ValueError(f'{cod} {periodo:%m/%Y} es posterior al corte ({corte:%m/%Y}).')
    # Los pagos de la carga inicial se rehacen con lo que diga la fila.
    dep.pagos.filter(origen=Cargo.Origen.MIGRACION).delete()
    existentes = {(c.concepto.codigo, c.periodo): c
                  for c in dep.cargos.select_related('concepto').prefetch_related('imputaciones')}
    creados = actualizados = borrados = 0
    for clave, monto in deseados.items():
        c = existentes.get(clave)
        if c and c.origen != Cargo.Origen.MIGRACION:
            raise ValueError(f'Ya existe un cargo del sistema {clave[0]} {clave[1]:%m/%Y} '
                             f'(S/ {c.monto:.2f}); no se sobrescribe.')
        if c:
            if c.monto != monto:
                if monto < c.total_imputado:
                    raise ValueError(f'{clave[0]} {clave[1]:%m/%Y}: ya tiene pagos por '
                                     f'S/ {c.total_imputado:.2f}, no puede bajar a S/ {monto:.2f}.')
                c.monto = monto
                c.save(update_fields=['monto'])
                actualizados += 1
            continue
        existentes[clave] = Cargo.objects.create(
            departamento=dep, concepto=_concepto(clave[0]), periodo=clave[1], monto=monto,
            origen=Cargo.Origen.MIGRACION, descripcion='Carga inicial (migración)',
            fecha_vencimiento=fecha_vencimiento(clave[1]))
        creados += 1
    for clave, c in list(existentes.items()):
        if c.origen == Cargo.Origen.MIGRACION and clave not in deseados:
            if c.total_imputado > 0:
                raise ValueError(f'{clave[0]} {clave[1]:%m/%Y} tiene pagos; no se puede quitar.')
            c.delete()
            borrados += 1

    cobradas = {k: redondear(v) for k, v in r.get('cobradas', {}).items() if v}
    if cobradas and fecha_cobro is None:
        raise ValueError('Falta la fecha de cobro de las deudas ya cobradas.')
    for clave, monto in sorted(cobradas.items()):
        cargo = existentes[(Concepto.MANTENIMIENTO, _periodo(clave))]
        libre = cargo.monto - sum((i.monto for i in cargo.imputaciones.all()), CERO)
        if monto > libre:
            raise ValueError(f'Cobrado {clave[5:]}/{clave[:4]} S/ {monto:.2f} supera lo que queda '
                             f'de ese mes (S/ {libre:.2f}).')
        pago = Pago.objects.create(departamento=dep, monto=monto, fecha_pago=fecha_cobro,
                                   concepto=cargo.concepto, metodo='Carga inicial',
                                   nota=f'Deuda {clave[5:]}/{clave[:4]} ya cobrada (carga inicial)',
                                   origen=Cargo.Origen.MIGRACION)
        Imputacion.objects.create(pago=pago, cargo=cargo, monto=monto)
    return creados, actualizados, borrados
