"""Lógica de negocio del cálculo y la generación de recibos."""
import unicodedata
from datetime import date
from decimal import InvalidOperation, ROUND_HALF_UP, Decimal

from django.db import transaction
from django.db.models import Sum
from django.utils import timezone

from .models import (
    Cargo, Concepto, Departamento, GastoEdificio, GastoFijo, Imputacion,
    LecturaMedidor, Pago, TarifaAgua,
)

CENTIMOS = Decimal('0.01')


def _redondear(valor):
    return Decimal(valor).quantize(CENTIMOS, rounding=ROUND_HALF_UP)


def _fecha_vencimiento(periodo, dia=5):
    """Día `dia` del mes siguiente al periodo (por defecto, el 5, como el Excel)."""
    anio, mes = periodo.year, periodo.month
    if mes == 12:
        anio, mes = anio + 1, 1
    else:
        mes += 1
    return date(anio, mes, dia)


# ---------------------------------------------------------------- Agua

def costo_agua(m3):
    """Costo del agua por tramos acumulativos según los tramos activos.

    Cada m³ se cobra a la tarifa del tramo en que cae (0–20 a 3.93, 20–50 a 5.52,
    etc.). El tramo final (sin tope) cubre todo el excedente.
    """
    m3 = Decimal(m3)
    if m3 <= 0:
        return _redondear('0')

    tramos = list(TarifaAgua.objects.filter(activo=True))
    tramos.sort(key=lambda t: (t.consumo_limite is None, t.consumo_limite or Decimal('0')))

    total = Decimal('0')
    piso = Decimal('0')
    for tramo in tramos:
        if m3 <= piso:
            break
        tope = tramo.consumo_limite if tramo.consumo_limite is not None else m3
        cantidad = min(m3, tope) - piso
        if cantidad > 0:
            total += cantidad * tramo.tarifa
        piso = tope
        if tramo.consumo_limite is None:
            break
    return _redondear(total)


def ultima_lectura(departamento, antes_de=None):
    qs = LecturaMedidor.objects.filter(departamento=departamento)
    if antes_de is not None:
        qs = qs.filter(periodo__lt=antes_de)
    return qs.order_by('-periodo').first()


@transaction.atomic
def registrar_lectura(departamento, periodo, lectura_actual, lectura_anterior=None):
    """Registra la lectura del mes. Arrastra la lectura anterior si no se pasa."""
    if lectura_anterior is None:
        previa = ultima_lectura(departamento, antes_de=periodo)
        lectura_anterior = previa.lectura_actual if previa else Decimal('0')
    lectura, _ = LecturaMedidor.objects.update_or_create(
        departamento=departamento, periodo=periodo,
        defaults={'lectura_anterior': lectura_anterior, 'lectura_actual': lectura_actual},
    )
    return lectura


# ---------------------------------------------- Prorrateo de gastos

def _por_depto(compartido_total, directo_total, n):
    """Monto por depto: la parte compartida se divide entre n; la parte
    'por departamento' se cobra completa a cada uno."""
    compartido = compartido_total / n if n else Decimal('0')
    return _redondear(compartido + directo_total)


def prorrateo_mes(periodo):
    """Mantenimiento y área común por departamento, combinando gastos fijos
    (recurrentes) y gastos adicionales del mes (dinámicos).

    Cada gasto puede ser compartido (se divide entre los deptos) o por
    departamento (cada depto paga el monto completo).
    """
    GENERAL = GastoEdificio.Categoria.GENERAL
    AGUA_COMUN = GastoEdificio.Categoria.AGUA_COMUN
    vacio = {'n': 0, 'mantenimiento': Decimal('0.00'), 'area_comun': Decimal('0.00'),
             'total_general': Decimal('0.00'), 'total_agua_comun': Decimal('0.00')}

    n = Departamento.objects.filter(activo=True).count()
    if n == 0:
        return vacio

    fijos = GastoFijo.objects.filter(activo=True)
    # Solo gastos del mes NO dirigidos (los personalizados se aplican por depto
    # en la generación, no en el prorrateo uniforme).
    mes = GastoEdificio.objects.filter(periodo=periodo, departamentos__isnull=True)

    def suma(categoria, compartido):
        f = fijos.filter(categoria=categoria, compartido=compartido).aggregate(
            s=Sum('monto'))['s'] or Decimal('0.00')
        m = mes.filter(categoria=categoria, compartido=compartido).aggregate(
            s=Sum('monto'))['s'] or Decimal('0.00')
        return f + m

    def suma_fijo(categoria, compartido):
        return fijos.filter(categoria=categoria, compartido=compartido).aggregate(
            s=Sum('monto'))['s'] or Decimal('0.00')

    def suma_mes(categoria, compartido):
        return mes.filter(categoria=categoria, compartido=compartido).aggregate(
            s=Sum('monto'))['s'] or Decimal('0.00')

    # Mantenimiento = gastos FIJOS (recurrentes). Adicionales = gastos del MES.
    mantenimiento = _por_depto(suma_fijo(GENERAL, True), suma_fijo(GENERAL, False), n)
    adicionales = _por_depto(suma_mes(GENERAL, True), suma_mes(GENERAL, False), n)
    # Área común combina fijos + del mes.
    area_comun = _por_depto(suma(AGUA_COMUN, True), suma(AGUA_COMUN, False), n)

    return {
        'n': n,
        'total_general': suma(GENERAL, True) + suma(GENERAL, False),
        'total_agua_comun': suma(AGUA_COMUN, True) + suma(AGUA_COMUN, False),
        'mantenimiento': mantenimiento,
        'adicionales': adicionales,
        'area_comun': area_comun,
    }


def desglose_gastos_mes(periodo):
    """Desglose itemizado de gastos del mes (fijos + del mes) para el balance.

    Por cada gasto devuelve su total (columna 'TORRE D' del Excel) y su aporte
    por departamento (columna 'TOTAL'), agrupado por categoría.
    """
    n = Departamento.objects.filter(activo=True).count() or 1

    def item(g, adicional, objetivos=0):
        divisor = objetivos if objetivos else n
        share = (g.monto / divisor) if g.compartido else g.monto
        return {'descripcion': g.descripcion, 'total': g.monto,
                'por_depto': _redondear(share), 'compartido': g.compartido,
                'adicional': adicional, 'objetivos': objetivos}

    general, agua_comun = [], []
    for g in GastoFijo.objects.filter(activo=True):
        (agua_comun if g.categoria == GastoEdificio.Categoria.AGUA_COMUN else general).append(
            item(g, False))
    for g in GastoEdificio.objects.filter(periodo=periodo).prefetch_related('departamentos'):
        objetivos = g.departamentos.count()
        (agua_comun if g.categoria == GastoEdificio.Categoria.AGUA_COMUN else general).append(
            item(g, True, objetivos))

    return {'n': n, 'general': general, 'agua_comun': agua_comun}


def desglose_adicionales(periodo):
    """Lista de gastos adicionales del mes con su aporte por departamento.

    Sirve para el '?' que muestra qué compone la columna de adicionales.
    """
    n = Departamento.objects.filter(activo=True).count() or 1
    items = []
    for g in GastoEdificio.objects.filter(periodo=periodo,
                                          categoria=GastoEdificio.Categoria.GENERAL):
        aporte = (g.monto / n) if g.compartido else g.monto
        items.append({'descripcion': g.descripcion, 'monto': _redondear(aporte),
                      'compartido': g.compartido})
    return items


# ---------------------------------------------- Conceptos de deuda

# (código, nombre, orden, es_recurrente)
CONCEPTOS_ESTANDAR = [
    ('MANTENIMIENTO', 'Mantenimiento', 10, True),
    ('AGUA', 'Agua', 20, True),
    ('AREA_COMUN', 'Área común', 30, True),
    ('COCHERA', 'Cochera', 40, True),
    ('GASTO_ADIC', 'Gastos adicionales', 50, True),
    ('PUERTAS', 'Puertas', 60, False),
    ('CAMARA', 'Cámaras', 70, False),
    ('ASCENSORES', 'Ascensores', 80, False),
    ('CUOTA', 'Cuota extraordinaria', 90, False),
    ('MULTA', 'Multa', 100, False),
    ('MORA', 'Mora', 110, False),
    ('OTRO', 'Otro', 120, False),
]

# Conceptos recurrentes → de dónde sale su monto por depto.
_RECURRENTES = ['MANTENIMIENTO', 'AGUA', 'AREA_COMUN', 'COCHERA', 'GASTO_ADIC']


def sembrar_conceptos():
    """Crea los conceptos estándar si faltan (idempotente)."""
    for codigo, nombre, orden, recurrente in CONCEPTOS_ESTANDAR:
        Concepto.objects.get_or_create(
            codigo=codigo,
            defaults={'nombre': nombre, 'orden': orden, 'es_recurrente': recurrente})


# ---------------------------------------------- Generación de cargos

@transaction.atomic
def generar_cargos_mes(periodo, dia_vencimiento=5):
    """Genera (o actualiza) los cargos recurrentes del mes por cada depto activo.

    Un cargo por concepto (mantenimiento, agua, área común, cochera, adicionales).
    Idempotente: no duplica; no toca cargos que ya recibieron pagos.
    """
    sembrar_conceptos()
    datos = prorrateo_mes(periodo)
    vencimiento = _fecha_vencimiento(periodo, dia_vencimiento)
    conceptos = {c.codigo: c for c in Concepto.objects.filter(codigo__in=_RECURRENTES)}
    resumen = {'creados': 0, 'actualizados': 0, 'omitidos_pagados': 0, 'n': datos['n']}

    # Gastos personalizados del mes: aporte extra por depto seleccionado.
    extra_por_dep = {}
    dirigidos = (GastoEdificio.objects.filter(periodo=periodo)
                 .filter(departamentos__isnull=False).distinct().prefetch_related('departamentos'))
    for g in dirigidos:
        objetivos = list(g.departamentos.all())
        if not objetivos:
            continue
        aporte = (g.monto / len(objetivos)) if g.compartido else g.monto
        for d in objetivos:
            extra_por_dep[d.id] = extra_por_dep.get(d.id, Decimal('0.00')) + aporte

    for dep in Departamento.objects.filter(activo=True):
        lectura = LecturaMedidor.objects.filter(departamento=dep, periodo=periodo).first()
        m3 = lectura.m3_consumidos if lectura else Decimal('0.00')

        montos = {
            'MANTENIMIENTO': datos['mantenimiento'],
            'AGUA': costo_agua(m3),
            'AREA_COMUN': datos['area_comun'],
            'COCHERA': _redondear(dep.cochera_monto or 0),
            'GASTO_ADIC': _redondear(datos['adicionales'] + extra_por_dep.get(dep.id, Decimal('0.00'))),
        }
        for codigo, monto in montos.items():
            concepto = conceptos[codigo]
            cargo = Cargo.objects.filter(departamento=dep, concepto=concepto, periodo=periodo).first()

            if monto <= 0:
                # Sin monto: si existe un cargo sin pagos, se elimina; si tiene pagos, se deja.
                if cargo and cargo.total_imputado == 0:
                    cargo.delete()
                continue

            if cargo is None:
                Cargo.objects.create(departamento=dep, concepto=concepto, periodo=periodo,
                                     monto=monto, fecha_vencimiento=vencimiento)
                resumen['creados'] += 1
            elif cargo.total_imputado > 0:
                resumen['omitidos_pagados'] += 1
            else:
                cargo.monto = monto
                cargo.fecha_vencimiento = vencimiento
                cargo.save(update_fields=['monto', 'fecha_vencimiento'])
                resumen['actualizados'] += 1

    return resumen


@transaction.atomic
def crear_cargo(departamento, concepto, periodo, monto, descripcion='', dia_vencimiento=5):
    """Crea o actualiza un cargo (puertas, cámara, cuota, multa, mora, etc.).

    Idempotente por (depto, concepto, periodo): si ya existe, reemplaza el monto
    y la descripción (evita duplicados y el error de unicidad).
    """
    cargo, _ = Cargo.objects.update_or_create(
        departamento=departamento, concepto=concepto, periodo=periodo,
        defaults={'monto': _redondear(monto), 'descripcion': descripcion,
                  'fecha_vencimiento': _fecha_vencimiento(periodo, dia_vencimiento)})
    return cargo


# ---------------------------------------------- Importación de Excel

# Encabezados aceptados (normalizados) -> campo interno.
_ENCABEZADOS = {
    'DEPARTAMENTO': 'numero', 'DPTO': 'numero', 'DEPTO': 'numero',
    'NUMERO': 'numero', 'N_DEPARTAMENTO': 'numero', 'N_DPTO': 'numero',
    'LECTURA_ANTERIOR': 'anterior', 'ANTERIOR': 'anterior',
    'LECTURA_ACTUAL': 'actual', 'ACTUAL': 'actual',
    'COCHERA': 'cochera', 'MONTO_COCHERA': 'cochera', 'N_COCHERA': 'cochera',
}


def _norm(valor):
    """Normaliza un encabezado: mayúsculas, sin acentos, con guion bajo."""
    texto = str(valor or '').strip().upper()
    texto = ''.join(c for c in unicodedata.normalize('NFD', texto)
                    if unicodedata.category(c) != 'Mn')
    return texto.replace(' ', '_')


def _a_decimal(valor):
    if valor is None or str(valor).strip() == '':
        return None
    try:
        return Decimal(str(valor).replace(',', '.'))
    except (InvalidOperation, ValueError):
        return None


@transaction.atomic
def importar_lecturas_excel(fileobj, periodo):
    """Lee un Excel con el formato del sistema y carga departamentos + lecturas.

    Formato esperado (primera hoja): una fila de encabezados con al menos la
    columna DEPARTAMENTO, y opcionalmente LECTURA_ANTERIOR, LECTURA_ACTUAL,
    COCHERA. Una fila por departamento.
    """
    import openpyxl

    wb = openpyxl.load_workbook(fileobj, data_only=True)
    ws = wb.active

    # Localiza la fila de encabezados (busca en las primeras filas).
    header_row, colmap = None, {}
    for r in range(1, min(ws.max_row, 15) + 1):
        fila = {}
        for c in range(1, min(ws.max_column, 30) + 1):
            campo = _ENCABEZADOS.get(_norm(ws.cell(row=r, column=c).value))
            if campo:
                fila[campo] = c
        if 'numero' in fila:
            header_row, colmap = r, fila
            break

    if header_row is None:
        raise ValueError('No se encontró la columna "DEPARTAMENTO" en la primera hoja.')

    resumen = {'departamentos': 0, 'lecturas': 0, 'errores': []}
    for r in range(header_row + 1, ws.max_row + 1):
        crudo = ws.cell(row=r, column=colmap['numero']).value
        if crudo is None or str(crudo).strip() == '':
            continue
        # Normaliza el número (101.0 -> "101"; "T6-301" -> "T6-301").
        num_dec = _a_decimal(crudo)
        numero = str(int(num_dec)) if (num_dec is not None and num_dec == num_dec.to_integral_value()) \
            else str(crudo).strip()

        defaults = {'activo': True}
        if 'cochera' in colmap:
            cochera = _a_decimal(ws.cell(row=r, column=colmap['cochera']).value)
            if cochera is not None:
                defaults['cochera_monto'] = cochera
        dep, _ = Departamento.objects.update_or_create(numero_domicilio=numero, defaults=defaults)
        resumen['departamentos'] += 1

        actual = _a_decimal(ws.cell(row=r, column=colmap['actual']).value) if 'actual' in colmap else None
        anterior = _a_decimal(ws.cell(row=r, column=colmap['anterior']).value) if 'anterior' in colmap else None
        if actual is not None:
            try:
                registrar_lectura(dep, periodo, lectura_actual=actual, lectura_anterior=anterior)
                resumen['lecturas'] += 1
            except Exception as e:  # pragma: no cover - fila con datos inconsistentes
                resumen['errores'].append(f'Depto {numero}: {e}')

    return resumen


@transaction.atomic
def registrar_pago(departamento, monto, concepto=None, metodo='', nota='', periodo=None):
    """Registra un pago y lo imputa a los cargos del depto (más antiguos primero).

    Si `concepto` está definido (ej. Mantenimiento), solo cancela cargos de ese
    concepto. Si `periodo` está definido, solo los de ese mes.
    """
    monto = _redondear(monto)
    pago = Pago.objects.create(departamento=departamento, monto=monto, metodo=metodo,
                               concepto=concepto, nota=nota)

    cargos = departamento.cargos.all().order_by('periodo', 'concepto__orden')
    if concepto is not None:
        cargos = cargos.filter(concepto=concepto)
    if periodo is not None:
        cargos = cargos.filter(periodo=periodo)

    restante = monto
    for cargo in cargos:
        if restante <= 0:
            break
        saldo = cargo.saldo
        if saldo <= 0:
            continue
        aplicar = min(restante, saldo)
        Imputacion.objects.create(pago=pago, cargo=cargo, monto=aplicar)
        restante -= aplicar

    return pago


# ---------------------------------------------- Consultas de deuda

def saldo_departamento(departamento):
    """Deuda total (saldo) de un departamento = Σ saldos de sus cargos."""
    return sum((c.saldo for c in departamento.cargos.all()), Decimal('0.00'))


def deuda_anterior(departamento, periodo):
    """Saldo acumulado de meses anteriores al `periodo` (deuda arrastrada)."""
    return sum((c.saldo for c in departamento.cargos.all() if c.periodo < periodo),
               Decimal('0.00'))


def saldos_por_concepto(departamento):
    """Dict {codigo_concepto: saldo} de un departamento (para el desglose DEUDORES)."""
    acum = {}
    for c in departamento.cargos.all():
        if c.saldo != 0:
            acum[c.concepto.codigo] = acum.get(c.concepto.codigo, Decimal('0.00')) + c.saldo
    return acum


def deudores(periodo=None):
    """Lista de departamentos con deuda > 0 y el desglose de su saldo.

    Devuelve dicts con: departamento, total, anterior, mes (saldo del `periodo`),
    por_concepto. Ordenado de mayor a menor deuda.
    """
    filas = []
    deps = (Departamento.objects.filter(activo=True)
            .prefetch_related('cargos__concepto', 'cargos__imputaciones'))
    for dep in deps:
        total = sum((c.saldo for c in dep.cargos.all()), Decimal('0.00'))
        if total <= 0:
            continue
        anterior = Decimal('0.00')
        mes = Decimal('0.00')
        por_concepto = {}
        for c in dep.cargos.all():
            if c.saldo == 0:
                continue
            por_concepto[c.concepto.codigo] = por_concepto.get(c.concepto.codigo, Decimal('0.00')) + c.saldo
            if periodo is not None:
                if c.periodo < periodo:
                    anterior += c.saldo
                elif c.periodo == periodo:
                    mes += c.saldo
        filas.append({'departamento': dep, 'total': total, 'anterior': anterior,
                      'mes': mes, 'por_concepto': por_concepto})
    filas.sort(key=lambda f: f['total'], reverse=True)
    return filas
