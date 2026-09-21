"""Importa teléfonos de propietario/inquilino y proveedores desde excel_data.xlsx.

Formato (hoja 1): tabla LISTAS con columnas B=# DEPART., C=PROPIETARIO (celular),
D=INQUILINO (celular); y a la derecha una lista de proveedores en col J (nombre)
y col L (teléfono).

Uso:
    python manage.py importar_contactos "C:\\...\\excel_data.xlsx"
"""
from django.core.management.base import BaseCommand, CommandError

from departamentos.models import Departamento, Proveedor

COL_DPTO, COL_PROP, COL_INQ = 2, 3, 4        # # DEPART., PROPIETARIO, INQUILINO
COL_PROV_NOMBRE, COL_PROV_TEL = 10, 12       # proveedor: nombre y teléfono


def _texto(v):
    return '' if v is None else str(v).strip()


def _telefono(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return ''.join(ch for ch in str(v) if ch.isdigit())


class Command(BaseCommand):
    help = 'Importa teléfonos de propietario/inquilino y proveedores desde el Excel.'

    def add_arguments(self, parser):
        parser.add_argument('ruta', help='Ruta al archivo excel_data.xlsx.')

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError('Falta openpyxl. Instala con: pip install openpyxl')

        try:
            wb = openpyxl.load_workbook(opts['ruta'], data_only=True)
        except FileNotFoundError:
            raise CommandError(f"No se encontró el archivo: {opts['ruta']}")
        ws = wb.active

        deptos = proveedores = 0
        for row in range(5, ws.max_row + 1):
            # --- Departamentos: teléfonos de propietario e inquilino ---
            num = ws.cell(row=row, column=COL_DPTO).value
            num_txt = _telefono(num) if isinstance(num, (int, float)) else _texto(num)
            if num_txt:
                cel_prop = _telefono(ws.cell(row=row, column=COL_PROP).value)
                cel_inq = _telefono(ws.cell(row=row, column=COL_INQ).value)
                defaults = {'activo': True}
                if cel_prop:
                    defaults['celular'] = cel_prop
                defaults['celular_inquilino'] = cel_inq
                defaults['tiene_inquilino'] = bool(cel_inq)
                Departamento.objects.update_or_create(numero_domicilio=num_txt, defaults=defaults)
                deptos += 1

            # --- Proveedores (columna de la derecha) ---
            nombre = _texto(ws.cell(row=row, column=COL_PROV_NOMBRE).value)
            tel = _telefono(ws.cell(row=row, column=COL_PROV_TEL).value)
            if nombre:
                Proveedor.objects.update_or_create(
                    nombre=nombre, defaults={'telefono': tel, 'activo': True})
                proveedores += 1

        # La fila 4 también trae un proveedor (TECNICO DE J.V.).
        nombre = _texto(ws.cell(row=4, column=COL_PROV_NOMBRE).value)
        tel = _telefono(ws.cell(row=4, column=COL_PROV_TEL).value)
        if nombre:
            Proveedor.objects.update_or_create(nombre=nombre, defaults={'telefono': tel, 'activo': True})
            proveedores += 1

        self.stdout.write(self.style.SUCCESS(
            f'{deptos} departamentos actualizados y {proveedores} proveedores importados.'))
