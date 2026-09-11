"""Importa datos iniciales desde el Excel de mantenimiento.

Carga los departamentos (con su cochera), las lecturas del mes, la tarifa de agua
por tramos y los gastos del edificio, para arrancar el sistema con datos reales.

Uso:
    python manage.py importar_excel "C:\\ruta\\MANTENIMIENTO JUNIO 2026.xlsx"
    python manage.py importar_excel "...xlsx" --periodo 2026-07
"""
from datetime import date
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError

from departamentos.models import Departamento, GastoFijo, TarifaAgua
from departamentos.servicios import registrar_lectura

# Columnas de la hoja AGUA (1-indexado): D=DPTO, E=lect anterior, F=lect actual, K=cochera
COL_DPTO, COL_ANT, COL_ACT, COL_COCHERA = 4, 5, 6, 11

# Tarifa de agua por tramos (hoja Hoja1).
TRAMOS = [('0-20 m³', Decimal('20'), Decimal('3.93')),
          ('20-50 m³', Decimal('50'), Decimal('5.52')),
          ('50+ m³', None, Decimal('5.52'))]

# Gastos del edificio (hoja DATOS) que definen el mantenimiento y el área común.
GASTOS_GENERAL = [
    ('Energía eléctrica (PLUZ)', Decimal('1084.36')),
    ('Energía eléctrica', Decimal('42.92')),
    ('Energía eléctrica', Decimal('3.28')),
    ('Limpieza', Decimal('1200')),
    ('Vigilancia', Decimal('3097.93')),
    ('Acopio', Decimal('26.25')),
    ('Ascensores', Decimal('200')),
    ('Materiales de limpieza', Decimal('100')),
    ('Gastos administrativos', Decimal('380')),
    ('Agua vigilancia', Decimal('9.4')),
]
GASTOS_AGUA_COMUN = [('Consumo de agua área común', Decimal('631.63'))]


def _dec(valor):
    if valor is None:
        return None
    try:
        return Decimal(str(valor))
    except (InvalidOperation, ValueError):
        return None


class Command(BaseCommand):
    help = 'Importa departamentos, lecturas, tarifa y gastos desde el Excel.'

    def add_arguments(self, parser):
        parser.add_argument('ruta', help='Ruta al archivo .xlsx.')
        parser.add_argument('--periodo', default='2026-07',
                            help='Periodo de las lecturas/gastos, YYYY-MM (default 2026-07).')

    def handle(self, *args, **opts):
        try:
            import openpyxl
        except ImportError:
            raise CommandError('Falta openpyxl. Instala con: pip install openpyxl')

        try:
            anio, mes = map(int, opts['periodo'].split('-'))
            periodo = date(anio, mes, 1)
        except (ValueError, TypeError):
            raise CommandError('Periodo inválido. Usa YYYY-MM.')

        try:
            wb = openpyxl.load_workbook(opts['ruta'], data_only=True)
        except FileNotFoundError:
            raise CommandError(f"No se encontró el archivo: {opts['ruta']}")

        if 'AGUA' not in wb.sheetnames:
            raise CommandError('El Excel no tiene la hoja "AGUA".')
        ws = wb['AGUA']

        # --- Tarifa de agua ---
        TarifaAgua.objects.all().delete()
        for nombre, limite, tarifa in TRAMOS:
            TarifaAgua.objects.create(nombre=nombre, consumo_limite=limite, tarifa=tarifa)
        self.stdout.write(self.style.SUCCESS(f'{len(TRAMOS)} tramos de tarifa creados.'))

        # --- Gastos fijos (recurrentes, compartidos entre todos los deptos) ---
        GastoFijo.objects.all().delete()
        for desc, monto in GASTOS_GENERAL:
            GastoFijo.objects.create(descripcion=desc, monto=monto, compartido=True,
                                     categoria='GENERAL')
        for desc, monto in GASTOS_AGUA_COMUN:
            GastoFijo.objects.create(descripcion=desc, monto=monto, compartido=True,
                                     categoria='AGUA_COMUN')
        self.stdout.write(self.style.SUCCESS(
            f'{len(GASTOS_GENERAL) + len(GASTOS_AGUA_COMUN)} gastos fijos creados.'))

        # --- Departamentos + lecturas ---
        deptos = lecturas = 0
        for row in range(3, ws.max_row + 1):
            dpto_raw = ws.cell(row=row, column=COL_DPTO).value
            dpto = _dec(dpto_raw)
            if dpto is None:
                continue
            numero = str(int(dpto))
            cochera = _dec(ws.cell(row=row, column=COL_COCHERA).value) or Decimal('0')

            dep, _ = Departamento.objects.update_or_create(
                numero_domicilio=numero,
                defaults={'cochera_monto': cochera, 'activo': True},
            )
            deptos += 1

            ant = _dec(ws.cell(row=row, column=COL_ANT).value)
            act = _dec(ws.cell(row=row, column=COL_ACT).value)
            if ant is not None and act is not None:
                registrar_lectura(dep, periodo, lectura_actual=act, lectura_anterior=ant)
                lecturas += 1

        self.stdout.write(self.style.SUCCESS(
            f'{deptos} departamentos y {lecturas} lecturas importadas para {periodo:%Y-%m}.'))
        self.stdout.write('Ahora corre: manage.py generar_recibos ' + opts['periodo'])
